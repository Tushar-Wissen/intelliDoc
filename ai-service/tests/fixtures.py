from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document


CONTRACT_HEADINGS = [
    "Master Services Agreement",
    "1. Definitions",
    "2. Services",
    "3. Fees and Payment",
    "4. Confidentiality",
    "5. Term and Termination",
    "6. Liability",
    "7. Indemnity",
    "8. Governing Law",
    "9. Signatures",
]


def contract_paragraph(heading: str) -> str:
    return (
        f"This paragraph under {heading} states the parties shall perform their "
        "obligations in good faith. Each sentence is complete and ordered for citations."
    )


def write_contract_pdf(path: Path, scanned_pages: set[int] | None = None) -> Path:
    scanned_pages = scanned_pages or set()
    doc = fitz.open()
    for index, heading in enumerate(CONTRACT_HEADINGS, start=1):
        page = doc.new_page()
        if index in scanned_pages:
            page.insert_text((72, 72), str(index), fontsize=8)
        else:
            page.insert_text((72, 72), heading, fontsize=18)
            y = 110
            for _ in range(8):
                page.insert_textbox(fitz.Rect(72, y, 540, y + 50), contract_paragraph(heading), fontsize=10)
                y += 55
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def write_corrupt_pdf(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"not-a-real-pdf")
    return path


def write_structured_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("Statement of Work", level=1)
    document.add_paragraph("This statement covers delivery of the extraction platform.")
    document.add_heading("Scope", level=2)
    document.add_paragraph("The supplier shall deliver parsing, OCR, and chunking.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Deliverable"
    table.cell(0, 1).text = "Due date"
    table.cell(1, 0).text = "Parsed corpus"
    table.cell(1, 1).text = "Week 2"
    document.add_heading("Acceptance", level=2)
    document.add_paragraph("Acceptance requires all story acceptance criteria to pass.")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path
