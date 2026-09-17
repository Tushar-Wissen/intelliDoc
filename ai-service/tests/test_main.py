from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def test_health_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "UP"
    assert data["service"] == "intellidoc-ai-service"
    assert data["model_loaded"] is True


def test_analyze_document_success():
    payload = {
        "document_id": "doc_test_100",
        "title": "Spring Boot and FastAPI Integration Guide.pdf",
        "content": "IntelliDoc demonstrates modern enterprise software architecture. The system achieved 99.9% uptime and recorded positive revenue growth of 25% across all regions.",
        "max_summary_length": 120
    }
    response = client.post("/api/v1/analyze", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["document_id"] == "doc_test_100"
    assert len(data["summary"]) > 0
    assert data["sentiment"] in ["POSITIVE", "NEUTRAL", "NEGATIVE"]
    assert isinstance(data["entities"], list)
    assert isinstance(data["key_topics"], list)
    assert data["confidence_score"] > 0


def test_analyze_document_empty_content():
    payload = {
        "document_id": "doc_test_101",
        "content": "   "
    }
    response = client.post("/api/v1/analyze", json=payload)
    assert response.status_code == 400


def test_selective_extraction_routes_only_weak_pages_to_ocr():
    payload = {
        "document_id": "doc_extract_100",
        "pages": [
            {"page_number": 2, "native_text": "", "ocr_text": "Scanned page content", "ocr_confidence": 0.91},
            {"page_number": 1, "native_text": "Native page contains reliable text."},
        ],
    }

    response = client.post("/api/v1/extract", json=payload)

    assert response.status_code == 200
    data = response.json()
    assert [page["page_number"] for page in data["pages"]] == [1, 2]
    assert [page["method"] for page in data["pages"]] == ["native", "ocr"]
    assert data["ocr_pages"] == [2]
    assert data["native_pages"] == [1]
    assert len(data["sections"]) == 2
    assert data["chunks"][0]["page_number"] == 1
    assert data["chunks"][0]["token_count"] <= 400
    assert data["combined_text"].index("Native page") < data["combined_text"].index("Scanned page")


def test_selective_extraction_rejects_documents_without_text():
    payload = {
        "document_id": "doc_extract_101",
        "pages": [{"page_number": 1, "native_text": "", "ocr_text": ""}],
    }

    response = client.post("/api/v1/extract", json=payload)

    assert response.status_code == 400


def test_file_extraction_reads_docx_text_layer():
    document = Document()
    document.add_heading("Contract Overview", level=1)
    document.add_paragraph("This DOCX contains reliable native text for extraction.")
    document_bytes = BytesIO()
    document.save(document_bytes)

    response = client.post(
        "/api/v1/extract/file",
        data={"document_id": "doc_file_100"},
        files={
            "file": (
                "contract.docx",
                document_bytes.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    data = response.json()
    assert data["native_pages"] == [1]
    assert data["ocr_pages"] == []
    assert "reliable native text" in data["combined_text"]
    assert data["sections"][0]["heading"] == "Contract Overview"
    assert data["chunks"][0]["section_id"] == data["sections"][0]["section_id"]


def test_file_extraction_rejects_unsupported_files():
    response = client.post(
        "/api/v1/extract/file",
        data={"document_id": "doc_file_101"},
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )

    assert response.status_code == 400


def test_qa_endpoint_success():
    payload = {
        "document_id": "doc_test_200",
        "context": "Spring Boot provides robust backend API services, while Python FastAPI handles AI model predictions and text analysis.",
        "question": "What handles AI model predictions?"
    }
    response = client.post("/api/v1/qa", json=payload)
    assert response.status_code == 200
    data = response.json()
    assert data["document_id"] == "doc_test_200"
    assert "FastAPI" in data["answer"] or "Python" in data["answer"]
    assert data["confidence"] > 0.5
    assert data["is_not_found"] is False


def test_qa_endpoint_returns_not_found_without_guessing():
    payload = {
        "document_id": "doc_test_201",
        "context": "The contract expires on 31 March 2027.",
        "question": "Who approved the budget?"
    }

    response = client.post("/api/v1/qa", json=payload)

    assert response.status_code == 200
    data = response.json()
    assert data["answer"] == "Not found in supplied document context."
    assert data["is_not_found"] is True
    assert data["citations"] == []
