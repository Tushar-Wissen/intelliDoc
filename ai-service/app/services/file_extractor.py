from io import BytesIO
from pathlib import Path

from app.schemas import ExtractionPageInput, ExtractionResponse
from app.services.extractor import SelectiveExtractor
from app.services.ocr_adapter import create_ocr_adapter


class DocumentFileExtractor:
    """Extracts PDF, DOCX, and image uploads into the selective page pipeline."""

    MAX_FILE_BYTES = 20 * 1024 * 1024
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}

    @classmethod
    def extract(cls, document_id: str, filename: str, content: bytes) -> ExtractionResponse:
        if not content:
            raise ValueError("Uploaded document cannot be empty.")
        if len(content) > cls.MAX_FILE_BYTES:
            raise ValueError("Uploaded document exceeds the 20 MB limit.")

        extension = Path(filename).suffix.lower()
        if extension not in cls.SUPPORTED_EXTENSIONS:
            raise ValueError("Unsupported file type. Use PDF, DOCX, PNG, JPG, or TIFF.")

        if extension == ".pdf":
            pages = cls._extract_pdf(content)
        elif extension == ".docx":
            pages = cls._extract_docx(content)
        else:
            pages = [cls._extract_image(content)]

        result = SelectiveExtractor.extract(document_id, pages)
        if not result.combined_text:
            raise ValueError("No extractable text was found in the uploaded document.")
        return result

    @staticmethod
    def _extract_pdf(content: bytes) -> list[ExtractionPageInput]:
        from pypdf import PdfReader
        import fitz

        reader = PdfReader(BytesIO(content))
        rendered_document = fitz.open(stream=content, filetype="pdf")
        pages = []
        for index, page in enumerate(reader.pages):
            native_text = page.extract_text() or ""
            ocr_text = None
            if not SelectiveExtractor.native_text_is_usable(native_text):
                ocr_text, ocr_confidence = DocumentFileExtractor._ocr_image(rendered_document[index].get_pixmap(dpi=200))
            else:
                ocr_confidence = None
            pages.append(ExtractionPageInput(page_number=index + 1, native_text=native_text,
                                              ocr_text=ocr_text, ocr_confidence=ocr_confidence))
        rendered_document.close()
        return pages

    @staticmethod
    def _extract_docx(content: bytes) -> list[ExtractionPageInput]:
        from docx import Document

        document = Document(BytesIO(content))
        blocks = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                blocks.append(paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        text = "\n".join(blocks).strip()
        return [ExtractionPageInput(page_number=1, native_text=text)]

    @staticmethod
    def _extract_image(content: bytes) -> list[ExtractionPageInput]:
        from PIL import Image

        image = Image.open(BytesIO(content))
        ocr_text, ocr_confidence = DocumentFileExtractor._ocr_image(image)
        return [ExtractionPageInput(page_number=1, ocr_text=ocr_text, ocr_confidence=ocr_confidence)]

    @staticmethod
    def _ocr_image(image) -> tuple[str, float | None]:
        try:
            return create_ocr_adapter().extract(image)
        except Exception as error:
            raise ValueError(
                "OCR is unavailable. Verify the configured OCR engine and upload a document with a text layer."
            ) from error