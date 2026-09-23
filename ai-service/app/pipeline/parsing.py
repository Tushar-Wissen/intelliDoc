"""Story 2.1 — PDF/DOCX text and structure extraction.

Primary parser: Docling. Fallback trigger (OPEN QUESTION §26, assumed here):
any Docling import/convert failure, then PyMuPDF for PDF and python-docx for DOCX.
"""

from __future__ import annotations

import io
import logging
import re
import tempfile
import uuid
from pathlib import Path
from typing import Protocol

from app.db.repository import (
    PageRecord,
    PipelineRepository,
    SectionRecord,
    get_repository,
)
from app.pipeline.types import ParsedDocument, ParsedPage, ParsedSection
from app.storage.minio_client import MinioReadError, MinioStorage, get_minio

logger = logging.getLogger(__name__)

_HEADING_NUMBER = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+(\S.*)$")


class ParseError(Exception):
    pass


class DocumentParser(Protocol):
    def parse(self, file_bytes: bytes, file_type: str) -> ParsedDocument:
        ...


def parse_document(
    document_id: uuid.UUID,
    *,
    repo: PipelineRepository | None = None,
    minio: MinioStorage | object | None = None,
    parser: DocumentParser | None = None,
) -> None:
    repo = repo or get_repository()
    minio = minio or get_minio()
    parser = parser or CascadingParser()

    document = repo.get_document(document_id)
    if document is None:
        raise ParseError("Document was not found.")

    try:
        file_bytes = minio.get_bytes(document.storage_path)
    except MinioReadError as exc:
        raise ParseError("The original file could not be read from storage.") from exc

    file_type = (document.file_type or "").lower().lstrip(".")
    try:
        parsed = parser.parse(file_bytes, file_type)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("The document could not be read. The file may be corrupt.") from exc

    if not parsed.pages:
        raise ParseError("The document could not be read. The file may be corrupt.")

    repo.clear_parse_artifacts(document_id)

    page_rows = [
        PageRecord(
            id=uuid.uuid4(),
            document_id=document_id,
            page_number=page.page_number,
            raw_text=page.raw_text,
            was_ocr=False,
            ocr_confidence=None,
        )
        for page in parsed.pages
    ]
    repo.replace_pages(document_id, page_rows)

    section_rows = _persistable_sections(document_id, parsed)
    repo.replace_sections(document_id, section_rows)
    logger.info(
        "Parsed documentId=%s pages=%s sections=%s",
        document_id,
        len(page_rows),
        len(section_rows),
    )


def _persistable_sections(
    document_id: uuid.UUID, parsed: ParsedDocument
) -> list[SectionRecord]:
    sections = parsed.sections or [
        ParsedSection(
            heading="Document",
            start_page=1,
            end_page=max(page.page_number for page in parsed.pages),
            body_text="\n".join(page.raw_text for page in parsed.pages),
        )
    ]
    ids = [uuid.uuid4() for _ in sections]
    rows: list[SectionRecord] = []
    for index, section in enumerate(sections):
        parent_id = None
        if section.parent_index is not None and 0 <= section.parent_index < index:
            parent_id = ids[section.parent_index]
        rows.append(
            SectionRecord(
                id=ids[index],
                document_id=document_id,
                parent_section_id=parent_id,
                heading=section.heading[:1024] if section.heading else "Document",
                start_page=section.start_page,
                end_page=section.end_page,
            )
        )
    return rows


class CascadingParser:
    """Docling first; PyMuPDF / python-docx if Docling fails to open or convert."""

    def parse(self, file_bytes: bytes, file_type: str) -> ParsedDocument:
        try:
            return DoclingParser().parse(file_bytes, file_type)
        except Exception as exc:
            logger.warning("Docling parse failed; using fallback. reason=%s", exc)
        if file_type == "pdf":
            return PyMuPdfParser().parse(file_bytes, file_type)
        if file_type == "docx":
            return DocxParser().parse(file_bytes, file_type)
        raise ParseError("The document could not be read. The file may be corrupt.")


class DoclingParser:
    def parse(self, file_bytes: bytes, file_type: str) -> ParsedDocument:
        try:
            from docling.document_converter import DocumentConverter
        except Exception as exc:  # pragma: no cover - optional heavy dependency
            raise ParseError("Docling is not available") from exc

        suffix = ".pdf" if file_type == "pdf" else ".docx"
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
            handle.write(file_bytes)
            temp_path = Path(handle.name)
        try:
            result = DocumentConverter().convert(str(temp_path))
            document = result.document
            pages = _pages_from_docling(document)
            sections = _sections_from_headings(pages)
            if not pages:
                raise ParseError("Docling returned no pages")
            return ParsedDocument(pages=pages, sections=sections, file_type=file_type)
        finally:
            temp_path.unlink(missing_ok=True)


def _pages_from_docling(document) -> list[ParsedPage]:
    pages: list[ParsedPage] = []
    export_pages = getattr(document, "pages", None)
    if export_pages:
        for index, page in enumerate(export_pages.values() if hasattr(export_pages, "values") else export_pages, start=1):
            text = ""
            if hasattr(page, "text"):
                text = page.text or ""
            elif hasattr(document, "export_to_markdown"):
                text = ""
            headings = [
                line.strip()
                for line in text.splitlines()
                if _looks_like_heading_line(line.strip())
            ]
            pages.append(ParsedPage(page_number=index, raw_text=text, headings=headings))
        if any(page.raw_text for page in pages):
            return pages

    markdown = ""
    if hasattr(document, "export_to_markdown"):
        markdown = document.export_to_markdown() or ""
    if not markdown:
        raise ParseError("Docling produced no text")
    chunks = re.split(r"\n(?=#{1,6}\s)", markdown)
    if len(chunks) == 1:
        return [ParsedPage(page_number=1, raw_text=markdown, headings=_heading_lines(markdown))]
    pages = []
    for index, chunk in enumerate(chunks, start=1):
        pages.append(
            ParsedPage(
                page_number=index,
                raw_text=chunk.strip(),
                headings=_heading_lines(chunk),
            )
        )
    return pages


class PyMuPdfParser:
    def parse(self, file_bytes: bytes, file_type: str) -> ParsedDocument:
        try:
            import fitz
        except Exception as exc:
            raise ParseError("PyMuPDF is not available") from exc
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception as exc:
            raise ParseError("The document could not be read. The file may be corrupt.") from exc
        pages: list[ParsedPage] = []
        try:
            for index, page in enumerate(doc, start=1):
                raw_text, headings = _pdf_page_text_and_headings(page)
                pages.append(ParsedPage(page_number=index, raw_text=raw_text, headings=headings))
        finally:
            doc.close()
        if not pages:
            raise ParseError("The document could not be read. The file may be corrupt.")
        return ParsedDocument(
            pages=pages,
            sections=_sections_from_headings(pages),
            file_type="pdf",
        )


def _pdf_page_text_and_headings(page) -> tuple[str, list[str]]:
    payload = page.get_text("dict")
    lines: list[str] = []
    headings: list[str] = []
    sizes: list[float] = []
    structured: list[tuple[str, float, bool]] = []
    for block in payload.get("blocks", []):
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = "".join(span.get("text", "") for span in spans).strip()
            if not text:
                continue
            size = max((span.get("size", 0.0) for span in spans), default=0.0)
            flags = 0
            for span in spans:
                flags |= int(span.get("flags", 0))
            bold = bool(flags & 2 ** 4) or bool(flags & 16)
            sizes.append(size)
            structured.append((text, size, bold))
            lines.append(text)
    median = sorted(sizes)[len(sizes) // 2] if sizes else 0.0
    for text, size, bold in structured:
        if _looks_like_heading_line(text) or (
            median and size >= median * 1.15 and len(text.split()) <= 12
        ) or (bold and len(text.split()) <= 10 and not text.endswith(".")):
            if text not in headings:
                headings.append(text)
    raw_text = "\n".join(lines)
    if not raw_text:
        raw_text = page.get_text() or ""
    return raw_text, headings


class DocxParser:
    def parse(self, file_bytes: bytes, file_type: str) -> ParsedDocument:
        try:
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except Exception as exc:
            raise ParseError("python-docx is not available") from exc

        try:
            document = Document(io.BytesIO(file_bytes))
        except Exception as exc:
            raise ParseError("The document could not be read. The file may be corrupt.") from exc

        body_lines: list[str] = []
        headings: list[tuple[str, int]] = []
        page_number = 1
        pages_acc: list[list[str]] = [[]]

        def current_lines() -> list[str]:
            return pages_acc[-1]

        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, document)
                if _is_page_break(paragraph):
                    if current_lines():
                        page_number += 1
                        pages_acc.append([])
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
                level = _heading_level(style_name, text)
                if level:
                    headings.append((text, level))
                current_lines().append(text)
                body_lines.append(text)
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                rendered = _render_table(table)
                if rendered:
                    current_lines().append(rendered)
                    body_lines.append(rendered)

        pages: list[ParsedPage] = []
        for index, lines in enumerate(pages_acc, start=1):
            if not lines and len(pages_acc) > 1:
                continue
            text = "\n".join(lines)
            page_headings = [line for line in lines if _looks_like_heading_line(line) or _heading_level("", line)]
            # Prefer style-based headings collected globally, filtered to this page
            page_headings = [h for h, _ in headings if h in lines] or page_headings
            pages.append(ParsedPage(page_number=index, raw_text=text, headings=page_headings))

        if not pages or not any(page.raw_text.strip() for page in pages):
            raise ParseError("The document could not be read. The file may be corrupt.")

        parsed_sections = _sections_from_styled_headings(pages, headings)
        return ParsedDocument(pages=pages, sections=parsed_sections, file_type="docx")


def _is_page_break(paragraph) -> bool:
    xml = paragraph._p.xml
    return "w:type=\"page\"" in xml or "w:type='page'" in xml


def _heading_level(style_name: str, text: str) -> int | None:
    match = re.match(r"Heading\s+(\d+)", style_name or "", re.IGNORECASE)
    if match:
        return int(match.group(1))
    numbered = _HEADING_NUMBER.match(text.strip())
    if numbered:
        return numbered.group(1).count(".") + 1
    if style_name.lower().startswith("title"):
        return 1
    return None


def _render_table(table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _heading_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if _looks_like_heading_line(line.strip())]


def _looks_like_heading_line(text: str) -> bool:
    if not text:
        return False
    if text.startswith("#"):
        return True
    if _HEADING_NUMBER.match(text):
        return True
    words = text.split()
    if len(words) > 12 or len(text) > 120:
        return False
    if text.endswith((".", ",", ";", ":")):
        return False
    return text[:1].isupper()


def _sections_from_headings(pages: list[ParsedPage]) -> list[ParsedSection]:
    collected: list[tuple[int, str, int]] = []
    for page in pages:
        for heading in page.headings:
            numbered = _HEADING_NUMBER.match(heading)
            level = numbered.group(1).count(".") + 1 if numbered else 1
            collected.append((page.page_number, heading, level))
    return _build_section_tree(pages, collected)


def _sections_from_styled_headings(
    pages: list[ParsedPage], headings: list[tuple[str, int]]
) -> list[ParsedSection]:
    collected: list[tuple[int, str, int]] = []
    heading_to_page: dict[str, int] = {}
    for page in pages:
        for heading in page.headings:
            heading_to_page.setdefault(heading, page.page_number)
    for heading, level in headings:
        page_number = heading_to_page.get(heading, pages[0].page_number if pages else 1)
        collected.append((page_number, heading, level))
    if not collected:
        return _sections_from_headings(pages)
    return _build_section_tree(pages, collected)


def _build_section_tree(
    pages: list[ParsedPage], collected: list[tuple[int, str, int]]
) -> list[ParsedSection]:
    if not collected:
        last_page = pages[-1].page_number if pages else 1
        return [
            ParsedSection(
                heading="Document",
                start_page=1,
                end_page=last_page,
                body_text="\n".join(page.raw_text for page in pages),
            )
        ]

    last_page = pages[-1].page_number
    sections: list[ParsedSection] = []
    stack: list[int] = []
    for start_page, heading, level in collected:
        parent_index = None
        while stack and sections[stack[-1]].level >= level:
            stack.pop()
        if stack:
            parent_index = stack[-1]
        sections.append(
            ParsedSection(
                heading=heading,
                start_page=start_page,
                end_page=last_page,
                parent_index=parent_index,
                level=level,
            )
        )
        stack.append(len(sections) - 1)

    for index, section in enumerate(sections):
        later_same_or_higher = [
            other.start_page
            for other in sections[index + 1 :]
            if other.level <= section.level
        ]
        if later_same_or_higher:
            section.end_page = max(section.start_page, later_same_or_higher[0])
        section.body_text = _section_body(pages, section)

    return sections


def _section_body(pages: list[ParsedPage], section: ParsedSection) -> str:
    parts: list[str] = []
    for page in pages:
        if section.start_page <= page.page_number <= section.end_page:
            parts.append(page.raw_text)
    return "\n".join(parts)
